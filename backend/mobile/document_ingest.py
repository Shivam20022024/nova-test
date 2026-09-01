"""
Document text extraction — PDF/DOCX/TXT/pasted text input modes for test
generation. Pure extraction, no interpretation: turns a file into plain
text; mobile/ai_gateway.py is what turns text into steps.
"""
from __future__ import annotations

import io


class DocumentExtractionError(Exception):
    pass


MAX_CHARS = 40_000  # keep prompt size sane; a requirement doc is not a novel


def _extract_pdf(raw: bytes) -> str:
    try:
        import pymupdf  # PyMuPDF
    except ImportError as exc:
        raise DocumentExtractionError("PDF extraction requires PyMuPDF (`pip install PyMuPDF`).") from exc
    try:
        doc = pymupdf.open(stream=raw, filetype="pdf")
        text = "\n".join(page.get_text() for page in doc)
    except Exception as exc:
        raise DocumentExtractionError(f"Could not parse PDF: {exc}") from exc
    return text


def _extract_docx(raw: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise DocumentExtractionError("DOCX extraction requires python-docx (`pip install python-docx`).") from exc
    try:
        document = docx.Document(io.BytesIO(raw))
        return "\n".join(p.text for p in document.paragraphs)
    except Exception as exc:
        raise DocumentExtractionError(f"Could not parse DOCX: {exc}") from exc


def _extract_xlsx(raw: bytes) -> str:
    """Test cases / user stories kept in a spreadsheet — one row per case is
    the common real-world shape (Title, Steps, Expected Result, ... columns
    in any order/naming, since this is free-text extraction, not the strict
    CSV importer in test_import.py). Renders every sheet as a readable table
    so the AI generator gets the same row/column structure a human reading
    the sheet would see, not a flattened word-soup."""
    try:
        import openpyxl
    except ImportError as exc:
        raise DocumentExtractionError("XLSX extraction requires openpyxl (`pip install openpyxl`).") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise DocumentExtractionError(f"Could not parse XLSX: {exc}") from exc

    sections = []
    try:
        for sheet in wb.worksheets:
            rows_text = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if not any(cells):
                    continue  # skip fully blank rows — common in exported sheets
                rows_text.append(" | ".join(cells))
            if rows_text:
                sections.append(f"Sheet: {sheet.title}\n" + "\n".join(rows_text))
    finally:
        wb.close()

    if not sections:
        raise DocumentExtractionError("Spreadsheet contained no non-empty rows.")
    return "\n\n".join(sections)


def _extract_txt(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("Could not decode text file (tried utf-8/latin-1).")


def extract_text(filename: str, raw: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = _extract_pdf(raw)
    elif lower.endswith(".docx"):
        text = _extract_docx(raw)
    elif lower.endswith(".xlsx"):
        text = _extract_xlsx(raw)
    elif lower.endswith(".xls"):
        raise DocumentExtractionError(
            "Legacy .xls format isn't supported — please re-save as .xlsx (File > Save As in Excel) and re-upload."
        )
    elif lower.endswith((".txt", ".md", ".csv")):
        text = _extract_txt(raw)
    else:
        raise DocumentExtractionError(f"Unsupported document type: {filename} (use .pdf, .docx, .xlsx, or .txt).")

    text = text.strip()
    if not text:
        raise DocumentExtractionError("Document contained no extractable text.")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n\n[...truncated...]"
    return text
